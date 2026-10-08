"""
Daily Dispatch Report Generator
--------------------------------
Reads the raw dispatch export copy-pasted from the ERP (tab-separated text,
exactly as it comes when you select the ERP grid and copy it) and produces
one polished Excel file with:

  1. The original detail table (Sr.#, Location, OB User:, OB Date:, DO#, DN#,
     DN/Dispatch Date, Working Area:, Party Name, Vehicle#, Bilty#, Freight:,
     Remarks, Item Name:, Case Count, Qty:, Total) -- formatted (wrap text,
     borders, merged Freight/Remarks/Total per DO# group, group fields
     repeated on every item row), with a live SUM formula for each group's
     Total and a Grand Total row.

  2. A pivot-style summary below it (Party Name / Item Name: / Vehicle# /
     Freight: / Total) with Party Name repeated on every row (matching
     Excel's "Repeat All Item Labels" pivot option, subtotals off), sorted
     by Party then Item, with a Grand Total row -- this mirrors the manual
     PivotTable the office currently builds by hand.

Usage (from the command line):
    python generate_dispatch_report.py raw_dispatch.txt Dispatch_Report_22-Sep-2026.xlsx "22-Sep-2026"

Or import parse_raw_dispatch() / build_report() directly from app.py.
"""

import sys
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GROUP_FIELD_NAMES = [
    "sr", "location", "ob_user", "ob_date", "do", "dn", "dispatch_date",
    "area", "party", "vehicle", "bilty", "freight", "remarks",
]

SKIP_FIRST_TOKENS = {"Sr.#", "Item Name:", "Total"}


def _is_group_start(fields):
    """A group-start line's first cell is the Sr.# -- a plain integer, and the
    line carries all the group columns (so a lone '1350' total is not one)."""
    first = fields[0].strip()
    return bool(re.fullmatch(r"\d+", first)) and len(fields) >= 10


def _is_number(s):
    return bool(re.fullmatch(r"[\d,]+(\.\d+)?", (s or "").strip()))


def _to_float(s):
    s = (s or "").strip().replace(",", "")
    if s == "" or s.lower() == "blank":
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_raw_dispatch(raw_text):
    """Parse tab-separated ERP dispatch text into a list of group dicts:
    {sr, location, ob_user, ob_date, do, dn, dispatch_date, area, party,
     vehicle, bilty, freight, remarks, items: [(item_name, case_count, qty), ...]}
    """
    groups = []
    current = None

    for line in raw_text.splitlines():
        if not line.strip():
            continue
        first_token = line.split("\t", 1)[0].strip()
        if first_token.startswith("Date:"):
            continue
        if first_token in SKIP_FIRST_TOKENS:
            continue

        fields = line.split("\t")

        if _is_group_start(fields):
            if current is not None:
                groups.append(current)
            vals = (fields + [""] * len(GROUP_FIELD_NAMES))[: len(GROUP_FIELD_NAMES)]
            current = dict(zip(GROUP_FIELD_NAMES, [v.strip() for v in vals]))
            current["sr"] = int(current["sr"]) if current["sr"].isdigit() else current["sr"]
            current["freight"] = _to_float(current["freight"])
            current["bilty"] = current["bilty"]
            current["items"] = []
            # Occasionally the first item rides on the same line if there was
            # no trailing blank cell -- handle both shapes defensively.
            if len(fields) > len(GROUP_FIELD_NAMES) and fields[len(GROUP_FIELD_NAMES)].strip():
                rest = fields[len(GROUP_FIELD_NAMES):]
                name = rest[0].strip()
                case_count = _to_float(rest[1]) if len(rest) > 1 else 0.0
                qty = _to_float(rest[2]) if len(rest) > 2 else 0.0
                if name:
                    current["items"].append((name, case_count, qty))
        else:
            # Item row: Item Name  [Case Count]  [Qty]  [Group Total -- last row only]
            if current is None:
                continue
            name = fields[0].strip()
            if not name:
                continue
            # A line that is just a number is the group's Total sitting on
            # its own line (not an item) -- ignore it.
            if _is_number(name):
                continue
            case_count = _to_float(fields[1]) if len(fields) > 1 else 0.0
            qty = _to_float(fields[2]) if len(fields) > 2 else 0.0
            current["items"].append((name, case_count, qty))

    if current is not None:
        groups.append(current)

    # Drop any group that ended up with no item rows (defensive)
    groups = [g for g in groups if g["items"]]
    return groups


# ---------------------------------------------------------------- styling --
ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
ROW_FILL = PatternFill("solid", fgColor="F6F8FF")
TOT_FILL = PatternFill("solid", fgColor="FFF2CC")
PIV_HDR_FILL = PatternFill("solid", fgColor="1F4E79")
GT_FILL = PatternFill("solid", fgColor="D9E1F2")

THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="404040")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_report(groups, report_date, out_path, sheet_name=None):
    """groups: output of parse_raw_dispatch(). report_date: display string,
    e.g. '22-Sep-2026'. Writes the workbook to out_path."""

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or report_date)[:31]

    headers = ["Sr.#", "Location", "OB User:", "OB Date:", "DO#", "DN#", "DN/Dispatch Date",
               "Working Area:", "Party Name", "Vehicle#", "Bilty#", "Freight:", "Remarks",
               "Item Name:", "Case Count", "Qty:", "Total"]
    NC = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
    c = ws.cell(row=1, column=1, value=f"Date: {report_date} to {report_date}")
    c.font = Font(name=ARIAL, size=9, bold=True)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border = Border(left=MED, right=MED, top=MED, bottom=MED)

    for i, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(name=ARIAL, size=9, bold=True, color="2B2B2B")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX

    r = 3
    first_data_row = r
    group_rows = []
    for g in groups:
        start = r
        for name, case_count, qty in g["items"]:
            ws.cell(row=r, column=1, value=g["sr"])
            ws.cell(row=r, column=2, value=g["location"])
            ws.cell(row=r, column=3, value=g["ob_user"])
            ws.cell(row=r, column=4, value=g["ob_date"])
            ws.cell(row=r, column=5, value=g["do"])
            ws.cell(row=r, column=6, value=g["dn"])
            ws.cell(row=r, column=7, value=g["dispatch_date"])
            ws.cell(row=r, column=8, value=g["area"])
            ws.cell(row=r, column=9, value=g["party"])
            ws.cell(row=r, column=10, value=g["vehicle"])
            ws.cell(row=r, column=11, value=g["bilty"])
            ws.cell(row=r, column=14, value=name)
            ws.cell(row=r, column=15, value=case_count)
            ws.cell(row=r, column=16, value=qty)
            r += 1
        end = r - 1
        ws.cell(row=start, column=12, value=g["freight"])
        ws.cell(row=start, column=13, value=g["remarks"])
        ws.cell(row=start, column=17, value=f"=SUM(P{start}:P{end})")
        if end > start:
            for col in (12, 13, 17):
                ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
        group_rows.append((g, start, end))

    last_data_row = r - 1
    for rr in range(first_data_row, last_data_row + 1):
        for cc in range(1, NC + 1):
            cell = ws.cell(row=rr, column=cc)
            cell.font = Font(name=ARIAL, size=9)
            cell.fill = ROW_FILL
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="left" if cc in (9, 13, 14) else "center",
                vertical="center", wrap_text=True)
            if cc in (11, 12, 15, 16, 17):
                cell.number_format = "#,##0"
        ws.cell(row=rr, column=17).font = Font(name=ARIAL, size=9, bold=True)
        ws.cell(row=rr, column=17).fill = TOT_FILL
        ws.row_dimensions[rr].height = 15

    tr = last_data_row + 1
    ws.cell(row=tr, column=1, value="Total")
    ws.cell(row=tr, column=17, value=f"=SUM(P{first_data_row}:P{last_data_row})")
    for cc in range(1, NC + 1):
        cell = ws.cell(row=tr, column=cc)
        cell.font = Font(name=ARIAL, size=9, bold=True)
        cell.fill = HDR_FILL
        cell.border = Border(left=THIN, right=THIN, top=MED, bottom=MED)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cc == 17:
            cell.number_format = "#,##0"

    # ------------------------------------------------ pivot-style summary --
    pr = tr + 3
    lbl = ws.cell(row=pr, column=1, value="Sum of Qty:")
    lbl.font = Font(name=ARIAL, size=10, bold=True)

    ph = ["Party Name", "Item Name:", "Vehicle#", "Freight:", "Total"]
    hrow = pr + 1
    for i, h in enumerate(ph, 1):
        c = ws.cell(row=hrow, column=i, value=h)
        c.font = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")
        c.fill = PIV_HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX

    flat = []
    for g, start, end in group_rows:
        for idx, (name, case_count, qty) in enumerate(g["items"]):
            flat.append((g["party"], name, g["vehicle"], g["freight"], start + idx))
    flat.sort(key=lambda x: (x[0].lower(), x[1].lower()))

    pv = hrow + 1
    for party, item, vehicle, freight, srcrow in flat:
        ws.cell(row=pv, column=1, value=party)
        ws.cell(row=pv, column=2, value=item)
        ws.cell(row=pv, column=3, value=vehicle)
        ws.cell(row=pv, column=4, value=freight)
        ws.cell(row=pv, column=5, value=f"=P{srcrow}")
        for cc in range(1, 6):
            cell = ws.cell(row=pv, column=cc)
            cell.font = Font(name=ARIAL, size=9)
            cell.fill = ROW_FILL
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="left" if cc in (1, 2) else "center",
                vertical="center", wrap_text=True)
            if cc in (4, 5):
                cell.number_format = "#,##0"
        ws.row_dimensions[pv].height = 15
        pv += 1

    gt = pv
    ws.cell(row=gt, column=1, value="Grand Total")
    ws.cell(row=gt, column=5, value=f"=SUM(E{hrow + 1}:E{pv - 1})")
    for cc in range(1, 6):
        cell = ws.cell(row=gt, column=cc)
        cell.font = Font(name=ARIAL, size=9, bold=True)
        cell.fill = GT_FILL
        cell.border = Border(left=THIN, right=THIN, top=MED, bottom=MED)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cc == 5:
            cell.number_format = "#,##0"

    widths = [6, 20, 13, 11, 15, 14, 12, 11, 34, 11, 9, 9, 28, 30, 10, 9, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 28

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    in_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else "Dispatch_Report.xlsx"
    report_date = sys.argv[3] if len(sys.argv) > 3 else "Today"

    with open(in_path, "r", encoding="utf-8") as f:
        raw_text = f.read()

    groups = parse_raw_dispatch(raw_text)
    build_report(groups, report_date, out_path)
    print(f"Report written to {out_path} ({len(groups)} DO groups)")